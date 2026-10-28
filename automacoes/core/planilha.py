"""Planilhas de entrada (.xlsx) das telas do painel: leitura e modelo para baixar."""
import io
from dataclasses import dataclass
from pathlib import Path
from typing import NamedTuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .texto import normalizar_codigo_pesquisa, sem_acento_maiusculo

EXTENSOES = (".xlsx", ".xlsm")
COLUNAS_CODIGO = ("COD", "CODIGO", "COD_CNPJ", "CPF_CNPJ", "CNPJ", "CPF")  # apelidos aceitos para o código


class PlanilhaInvalida(ValueError):
    """A planilha não pôde ser usada; a mensagem vai para a tela."""


def conferir_extensao(nome_arquivo: str):
    extensao = Path(nome_arquivo or "").suffix.lower()
    if extensao == ".xls":
        raise PlanilhaInvalida(
            "Planilhas .xls (Excel 97-2003) não são aceitas. No Excel, use Salvar como > Pasta de Trabalho do Excel (.xlsx)."
        )
    if extensao not in EXTENSOES:
        raise PlanilhaInvalida("Use uma planilha .xlsx.")


def nome_coluna(valor) -> str:
    """' Ação cielo ' -> 'ACAO_CIELO'. Assim o cabeçalho não depende de acento, espaço ou maiúscula."""
    return sem_acento_maiusculo(valor).replace(" ", "_")


def texto_celula(valor) -> str:
    """Valor da célula como texto. Número inteiro sai sem '.0' (12345.0 -> '12345')."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def ler_linhas(arquivo) -> tuple[list[str], list[tuple[int, dict]]]:
    """Lê a primeira aba. A primeira linha é o cabeçalho.

    Devolve as colunas (já com nome_coluna) e, para cada linha com algum valor,
    (número da linha no Excel, {coluna: texto}).
    """
    try:
        livro = openpyxl.load_workbook(arquivo, read_only=True, data_only=True)
    except Exception as erro:
        raise PlanilhaInvalida(f"Não consegui abrir a planilha ({type(erro).__name__}). Confira se é um .xlsx válido.")
    try:
        linhas = livro.worksheets[0].iter_rows(values_only=True)
        cabecalho = [nome_coluna(valor) for valor in next(linhas, ())]
        if not any(cabecalho):
            raise PlanilhaInvalida("A planilha está vazia: a primeira linha deve ter os nomes das colunas.")
        dados = []
        for numero, valores in enumerate(linhas, start=2):
            linha = {coluna: texto_celula(valor) for coluna, valor in zip(cabecalho, valores) if coluna}
            if any(linha.values()):
                dados.append((numero, linha))
        return cabecalho, dados
    finally:
        livro.close()


def primeira_coluna(colunas: list[str], nomes: tuple[str, ...]) -> str | None:
    """A primeira de `nomes` que existe na planilha (para aceitar apelidos, como COD ou CODIGO)."""
    return next((nome for nome in nomes if nome in colunas), None)


def coluna_codigo(colunas: list[str]) -> str:
    coluna = primeira_coluna(colunas, COLUNAS_CODIGO)
    if not coluna:
        raise PlanilhaInvalida("Coluna COD não encontrada na primeira linha da planilha.")
    return coluna


class LeituraCodigos(NamedTuple):
    codigos: list[str]
    linhas_sem_codigo: int


def ler_codigos(arquivo, tipo_pesquisa: str) -> LeituraCodigos:
    """Só a coluna de código, já normalizada para a pesquisa. Linhas sem código são puladas."""
    colunas, linhas = ler_linhas(arquivo)
    coluna = coluna_codigo(colunas)
    codigos, sem_codigo = [], 0
    for _, linha in linhas:
        valor = linha.get(coluna, "")
        if valor:
            codigos.append(normalizar_codigo_pesquisa(valor, tipo_pesquisa))
        else:
            sem_codigo += 1
    if not codigos:
        raise PlanilhaInvalida(f"Nenhum código preenchido na coluna {coluna}.")
    return LeituraCodigos(codigos, sem_codigo)


# --- planilha modelo para baixar ----------------------------------------------
# A aba lida (a primeira) vem só com o cabeçalho: exemplos ficam na aba de instruções,
# porque um código de exemplo esquecido ali seria pesquisado (ou alterado) no portal.
LINHAS_MODELO = 2000  # até onde vão as listas de valores e as dicas

_FUNDO_CABECALHO = {
    "obrigatoria": PatternFill("solid", fgColor="1F3864"),
    "opcional": PatternFill("solid", fgColor="2E75B6"),
    "ignorada": PatternFill("solid", fgColor="8C9AB0"),
}
_BRANCO = Font(bold=True, color="FFFFFF")


@dataclass(frozen=True)
class ColunaModelo:
    nome: str
    largura: int
    dica: str                      # aparece no Excel ao clicar numa célula da coluna (até 255 caracteres)
    opcoes: tuple[str, ...] = ()   # se houver, vira lista de valores no Excel
    papel: str = "opcional"        # "obrigatoria", "opcional" ou "ignorada": só muda a cor do cabeçalho


def lista_de_opcoes(valores) -> str:
    """('A', 'B', 'C') -> 'A, B ou C'."""
    return ", ".join(valores[:-1]) + " ou " + valores[-1] if len(valores) > 1 else "".join(valores)


def montar_modelo(colunas: list[ColunaModelo], *, titulo: str, avisos: list[str],
                  regras: list[tuple], exemplo: list[tuple]) -> bytes:
    """Planilha .xlsx pronta para preencher: aba "Credenciados" (a lida pelo painel) e aba "Como preencher".

    `regras` são linhas (coluna, obrigatória, valores aceitos, se ficar vazia); `exemplo` segue a ordem de `colunas`.
    """
    livro = openpyxl.Workbook()
    aba = livro.active
    aba.title = "Credenciados"
    for indice, coluna in enumerate(colunas, start=1):
        letra = get_column_letter(indice)
        celula = aba.cell(row=1, column=indice, value=coluna.nome)
        celula.font = _BRANCO
        celula.alignment = Alignment(horizontal="center")
        celula.fill = _FUNDO_CABECALHO[coluna.papel]
        aba.column_dimensions[letra].width = coluna.largura
        aba.column_dimensions[letra].number_format = "@"  # texto: o Excel não come zeros à esquerda nem vira 1,23E+13

        if coluna.opcoes:
            validacao = DataValidation(
                type="list", formula1='"' + ",".join(coluna.opcoes) + '"', allow_blank=True,
                showErrorMessage=True, errorTitle=f"{coluna.nome} inválido"[:32],
                error=f"Use {lista_de_opcoes(coluna.opcoes)}, ou deixe a célula vazia.",
            )
        else:
            validacao = DataValidation(allow_blank=True)  # sem regra, só para mostrar a dica
        validacao.showInputMessage = True
        validacao.promptTitle = coluna.nome[:32]
        validacao.prompt = coluna.dica
        validacao.add(f"{letra}2:{letra}{LINHAS_MODELO}")
        aba.add_data_validation(validacao)
    aba.freeze_panes = "A2"

    _aba_instrucoes(livro.create_sheet("Como preencher"), colunas, titulo, avisos, regras, exemplo)
    saida = io.BytesIO()
    livro.save(saida)
    return saida.getvalue()


def _aba_instrucoes(aba, colunas, titulo, avisos, regras, exemplo):
    aba["A1"] = titulo
    aba["A1"].font = Font(bold=True, size=14, color="1F3864")
    linha = 2
    for aviso in ["Preencha a aba \"Credenciados\", uma linha por credenciado. Só a primeira aba é lida: "
                  "esta aqui é ignorada pelo robô.", *avisos]:
        aba.cell(row=linha, column=1, value=aviso)
        linha += 1

    linha += 1
    for indice, cabecalho in enumerate(("Coluna", "Obrigatória", "Valores aceitos", "Se ficar vazia"), start=1):
        celula = aba.cell(row=linha, column=indice, value=cabecalho)
        celula.font, celula.fill = _BRANCO, _FUNDO_CABECALHO["obrigatoria"]
    for valores in regras:
        linha += 1
        for indice, valor in enumerate(valores, start=1):
            aba.cell(row=linha, column=indice, value=valor)

    linha += 2
    aba.cell(row=linha, column=1, value="Exemplo (copie para a aba Credenciados só o que for real):").font = Font(bold=True)
    linha += 1
    for indice, coluna in enumerate(colunas, start=1):
        celula = aba.cell(row=linha, column=indice, value=coluna.nome)
        celula.font, celula.fill = _BRANCO, _FUNDO_CABECALHO["opcional"]
    for valores in exemplo:
        linha += 1
        for indice, valor in enumerate(valores, start=1):
            aba.cell(row=linha, column=indice, value=valor or None)

    for letra, largura in zip("ABCDEF", (16, 22, 30, 52, 16, 40)):
        aba.column_dimensions[letra].width = largura

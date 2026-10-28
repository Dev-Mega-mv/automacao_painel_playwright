"""Gravação das planilhas de saída (.xlsx) das automações."""
import os
import re
import tempfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_FINA = Side(style="thin")
BORDA = Border(left=_FINA, right=_FINA, top=_FINA, bottom=_FINA)
CENTRO = Alignment(horizontal="center", vertical="center")
ESQUERDA = Alignment(horizontal="left", vertical="center", wrap_text=True)


def cor(hex_rgb: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_rgb)


def valor_excel(valor):
    """Valor seguro para uma célula: tira caracteres de controle que o Excel recusa."""
    if valor is None:
        return ""
    if isinstance(valor, (int, float, bool, datetime)):
        return valor
    return re.sub(r"[\x00-\x08\x0B-\x0C\x0E-\x1F]", "", str(valor))


def salvar_xlsx(livro, caminho: Path):
    """Grava num arquivo temporário e troca no fim, para nunca deixar um .xlsx pela metade."""
    descritor, temporario = tempfile.mkstemp(suffix=".xlsx", dir=caminho.parent)
    os.close(descritor)
    try:
        livro.save(temporario)
        os.replace(temporario, caminho)
    finally:
        if os.path.exists(temporario):
            os.remove(temporario)


def salvar_tabela(caminho: Path, aba: str, colunas: list[str], linhas: list[dict], *, coluna_resultado: str = ""):
    """Uma aba com cabeçalho e uma linha por dicionário.

    Se `coluna_resultado` for informada, a célula dessa coluna fica vermelha quando o texto começa
    com ERRO ou AVISO, para os problemas saltarem aos olhos.
    """
    livro = openpyxl.Workbook()
    planilha = livro.active
    planilha.title = aba

    for indice, titulo in enumerate(colunas, start=1):
        celula = planilha.cell(row=1, column=indice, value=titulo)
        celula.fill = cor("1F3864")
        celula.font = Font(bold=True, color="FFFFFF", size=10)
        celula.alignment = CENTRO
        celula.border = BORDA

    for numero, dados in enumerate(linhas, start=2):
        for indice, titulo in enumerate(colunas, start=1):
            celula = planilha.cell(row=numero, column=indice, value=valor_excel(dados.get(titulo, "")))
            celula.border = BORDA
            celula.font = Font(size=10)
            celula.alignment = ESQUERDA
        resultado = str(dados.get(coluna_resultado, "")) if coluna_resultado else ""
        if resultado.startswith(("ERRO", "AVISO")):
            planilha.cell(row=numero, column=colunas.index(coluna_resultado) + 1).fill = cor("FFCCCC")

    for indice, titulo in enumerate(colunas, start=1):
        maior = max([len(titulo)] + [len(str(linha.get(titulo, ""))) for linha in linhas])
        planilha.column_dimensions[get_column_letter(indice)].width = min(max(12, maior + 2), 70)
    planilha.freeze_panes = "A2"
    salvar_xlsx(livro, caminho)

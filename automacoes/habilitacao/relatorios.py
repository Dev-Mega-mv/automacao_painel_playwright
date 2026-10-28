"""Planilhas da habilitação.

- Verificação: mesmo layout da versão Selenium (automacoes/cielo_cardse.py, _gerar_relatorio_xlsx).
- Ativação/desativação: uma linha por contrato com o antes e o depois de cada rede.
"""
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

from ..core.relatorio import BORDA, CENTRO, ESQUERDA, cor, salvar_tabela, salvar_xlsx, valor_excel


def gerar_relatorio_verificacao(dados: list[dict], caminho: Path):
    """dados = [{'cod': '123', 'produtos': [{'num', 'nome', 'cielo_sit', 'cielo_hab', 'cardse_sit', 'cardse_hab', 'obs'}]}]"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Verificação"
    colunas = [
        "#", "Produto",
        "CIELO — Situação Arquivo", "CIELO — Habilitado",
        "CARDSE — Situação Arquivo", "CARDSE — Habilitado",
        "Observação",
    ]

    ws.merge_cells("A1:G1")
    ws["A1"] = f"Relatório de Verificação — gerado em {datetime.now():%d/%m/%Y %H:%M:%S}"
    ws["A1"].fill = cor("1F3864")
    ws["A1"].font = Font(bold=True, color="FFFFFF", size=13)
    ws["A1"].alignment = CENTRO
    ws.row_dimensions[1].height = 28

    linha = 2
    for bloco in dados:
        ws.merge_cells(f"A{linha}:G{linha}")
        ws[f"A{linha}"] = valor_excel(f"  Credenciado: {bloco.get('cod', '—')}")
        ws[f"A{linha}"].fill = cor("2E75B6")
        ws[f"A{linha}"].font = Font(bold=True, color="FFFFFF", size=11)
        ws[f"A{linha}"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[linha].height = 22
        linha += 1

        for coluna, titulo in enumerate(colunas, start=1):
            celula = ws.cell(row=linha, column=coluna, value=titulo)
            celula.fill = cor("BDD7EE")
            celula.font = Font(bold=True, size=10)
            celula.alignment = CENTRO
            celula.border = BORDA
        linha += 1

        produtos = bloco.get("produtos", [])
        if not produtos:
            ws.merge_cells(f"A{linha}:G{linha}")
            ws[f"A{linha}"] = valor_excel(bloco.get("obs") or "Nenhum produto encontrado")
            ws[f"A{linha}"].alignment = CENTRO
            ws[f"A{linha}"].font = Font(italic=True, size=10, color="888888")
            linha += 1

        for produto in produtos:
            cielo_sit = (produto.get("cielo_sit") or "").strip().upper()
            cielo_hab = (produto.get("cielo_hab") or "").strip().upper()
            cardse_sit = (produto.get("cardse_sit") or "").strip().upper()
            cardse_hab = (produto.get("cardse_hab") or "").strip().upper()
            valores = [
                produto.get("num", ""), produto.get("nome", ""),
                cielo_sit or "—", cielo_hab or "—",
                cardse_sit or "—", cardse_hab or "—",
                produto.get("obs", ""),
            ]
            cores = [None, None, _cor_situacao(cielo_sit), _cor_habilitado(cielo_hab),
                     _cor_situacao(cardse_sit), _cor_habilitado(cardse_hab), None]
            for coluna, (valor, preenchimento) in enumerate(zip(valores, cores), start=1):
                celula = ws.cell(row=linha, column=coluna, value=valor_excel(valor))
                celula.font = Font(size=10)
                celula.border = BORDA
                celula.alignment = ESQUERDA if coluna in (2, 7) else CENTRO
                if preenchimento:
                    celula.fill = preenchimento
            linha += 1
        linha += 1

    for coluna, largura in zip("ABCDEFG", (6, 38, 24, 16, 24, 18, 28)):
        ws.column_dimensions[coluna].width = largura
    salvar_xlsx(wb, caminho)


def _cor_situacao(situacao: str):
    return cor("DDEBF7") if situacao and situacao != "—" else cor("F2F2F2")


def _cor_habilitado(habilitado: str):
    if habilitado == "SIM":
        return cor("C6EFCE")
    if habilitado in ("NÃO", "NAO"):
        return cor("FFCCCC")
    return cor("F2F2F2")


COLUNAS_HABILITACAO = [
    "Credenciado", "Contrato", "Produto", "Aprovado pelo robô",
    "CIELO antes", "CIELO depois", "CARDSE antes", "CARDSE depois", "Resultado",
]


def gerar_relatorio_habilitacao(linhas: list[dict], caminho: Path):
    """linhas = [{coluna de COLUNAS_HABILITACAO: valor}], uma por contrato processado."""
    salvar_tabela(caminho, "Habilitação", COLUNAS_HABILITACAO, linhas, coluna_resultado="Resultado")

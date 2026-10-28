"""Planilha da coleta: mesmo layout da versão Selenium (cielo_cardse.py, _gerar_relatorio_coleta_xlsx).

- Aba "Resumo": quando foi gerada, origem procurada, progresso e ocorrências.
- Aba "Coleta": uma linha por credenciado, com as colunas de COLUNAS.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ..core.relatorio import BORDA, cor, salvar_xlsx, valor_excel

COLUNAS = [
    "Codigo Infox/CNPJ", "Origem de Transacao", "Codigo Estabelecimento",
    "Data Cadastro", "Data Alteracao", "Situacao",
    "Contratos Totais", "Resumo Contratos", "Detalhes Contratos", "Observacao",
]
COLUNA_TOTAL = COLUNAS.index("Contratos Totais") + 1
COLUNA_RESUMO = COLUNAS.index("Resumo Contratos") + 1
COLUNA_OBSERVACAO = COLUNAS.index("Observacao") + 1

_TITULO = cor("0F3D5E")
_CABECALHO = cor("0F766E")
_ROTULO = cor("D9F99D")
_INFO = cor("F0FDFA")
_ERRO = cor("FEF2F2")
_OK = cor("ECFDF5")
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ESQUERDA = Alignment(horizontal="left", vertical="top", wrap_text=True)


def gravar_relatorio_coleta(caminho: Path, linhas: list[dict], resumo: list[tuple[str, object]], ocorrencias: list[str]):
    """`linhas` usa as chaves de COLUNAS; `resumo` são pares (campo, valor) da aba Resumo."""
    livro = openpyxl.Workbook()
    _aba_resumo(livro.active, resumo, ocorrencias)
    _aba_coleta(livro.create_sheet("Coleta"), linhas)
    salvar_xlsx(livro, caminho)


def _aba_resumo(aba, resumo, ocorrencias):
    aba.title = "Resumo"
    aba.merge_cells("A1:B1")
    aba["A1"] = "Resumo da Coleta de Codigo Estabelecimento"
    aba["A1"].fill = _TITULO
    aba["A1"].font = Font(bold=True, color="FFFFFF", size=12)
    aba["A1"].alignment = _CENTRO
    aba.row_dimensions[1].height = 26

    for linha, (campo, valor) in enumerate(resumo, start=3):
        rotulo = aba.cell(row=linha, column=1, value=valor_excel(campo))
        dado = aba.cell(row=linha, column=2, value=valor_excel(valor))
        for celula in (rotulo, dado):
            celula.border = BORDA
            celula.alignment = _ESQUERDA
        rotulo.fill = _ROTULO
        rotulo.font = Font(bold=True, size=10)
        dado.fill = _INFO

    if ocorrencias:
        linha = len(resumo) + 5
        aba.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
        titulo = aba.cell(row=linha, column=1, value="Ocorrencias registradas")
        titulo.fill = _TITULO
        titulo.font = Font(bold=True, color="FFFFFF", size=12)
        titulo.alignment = _ESQUERDA
        for linha, ocorrencia in enumerate(ocorrencias, start=linha + 1):
            aba.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=2)
            celula = aba.cell(row=linha, column=1, value=valor_excel(ocorrencia))
            celula.fill = _ERRO
            celula.border = BORDA
            celula.alignment = _ESQUERDA
    _ajustar_larguras(aba)


def _aba_coleta(aba, linhas):
    for coluna, titulo in enumerate(COLUNAS, start=1):
        celula = aba.cell(row=1, column=coluna, value=titulo)
        celula.fill = _CABECALHO
        celula.font = Font(bold=True, color="FFFFFF", size=10)
        celula.alignment = _CENTRO
        celula.border = BORDA

    for numero, dados in enumerate(linhas, start=2):
        tem_observacao = bool(dados.get("Observacao"))
        for coluna, titulo in enumerate(COLUNAS, start=1):
            celula = aba.cell(row=numero, column=coluna, value=valor_excel(dados.get(titulo, "")))
            celula.border = BORDA
            celula.alignment = _CENTRO if coluna == COLUNA_TOTAL else _ESQUERDA
            if numero % 2 == 0:
                celula.fill = _INFO
            if tem_observacao and coluna == COLUNA_OBSERVACAO:
                celula.fill = _ERRO
            elif coluna in (COLUNA_TOTAL, COLUNA_RESUMO) and not tem_observacao:
                celula.fill = _OK
        aba.row_dimensions[numero].height = 42 if "\n" in str(dados.get("Detalhes Contratos", "")) else 24

    _ajustar_larguras(aba)
    aba.freeze_panes = "A2"


def _ajustar_larguras(aba):
    """Largura pela linha de texto mais longa de cada coluna (entre 12 e 80)."""
    larguras = {}
    for linha in aba.iter_rows():
        for celula in linha:
            texto = "" if celula.value is None else str(celula.value)
            maior = max((len(parte) for parte in texto.splitlines()), default=0)
            larguras[celula.column] = max(larguras.get(celula.column, 0), min(maior + 2, 80))
    for coluna, largura in larguras.items():
        aba.column_dimensions[get_column_letter(coluna)].width = max(12, largura)
